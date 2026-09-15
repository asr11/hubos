#!/usr/bin/env python3
import sqlite3
import json
import os
import shutil
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DB_PATH = '/home/asr11/02_HUB_Dev/os.hub.co.il/server/data/os_hub.db'
BACKUPS_BASE = '/home/asr11/02_HUB_Dev/os.hub.co.il/backups'
MORNING_DIR = '/home/asr11/02_HUB_Dev/os.hub.co.il/server/data/morning'

def create_styled_sheet(wb, title, headers, rows, header_color="1E3A8A"):
    ws = wb.create_sheet(title=title)
    ws.views.sheetView[0].rightToLeft = True
    
    # Headers
    font_header = Font(name='Arial', size=11, bold=True, color='FFFFFF')
    fill_header = PatternFill(start_color=header_color, end_color=header_color, fill_type='solid')
    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    border_thin = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )
    
    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_thin
    
    # Rows
    font_body = Font(name='Arial', size=10)
    for row_data in rows:
        formatted_row = []
        for val in row_data:
            if isinstance(val, (dict, list)):
                formatted_row.append(json.dumps(val, ensure_ascii=False))
            elif val is None:
                formatted_row.append("")
            else:
                formatted_row.append(str(val))
        ws.append(formatted_row)
        
    for r in range(2, ws.max_row + 1):
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = font_body
            cell.border = border_thin
            cell.alignment = Alignment(vertical='center', wrap_text=False)
            
    # Auto column width
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)
        
    return ws

def run_backup():
    now = datetime.now()
    timestamp_str = now.strftime('%Y-%m-%d_%H-%M-%S')
    year = now.strftime('%Y')
    month = now.strftime('%m')
    day = now.strftime('%d')
    
    target_dir = os.path.join(BACKUPS_BASE, year, month, day, f"os_hub_backup_{timestamp_str}")
    excel_dir = os.path.join(target_dir, "master_excel")
    sqlite_dir = os.path.join(target_dir, "sqlite_db")
    json_dir = os.path.join(target_dir, "json_collections")
    morning_copy_dir = os.path.join(target_dir, "morning_csvs")
    
    os.makedirs(excel_dir, exist_ok=True)
    os.makedirs(sqlite_dir, exist_ok=True)
    os.makedirs(json_dir, exist_ok=True)
    os.makedirs(morning_copy_dir, exist_ok=True)
    
    # 1. Copy SQLite database safely
    db_backup_path = os.path.join(sqlite_dir, f"os_hub_{timestamp_str}.db")
    conn_src = sqlite3.connect(DB_PATH)
    conn_dst = sqlite3.connect(db_backup_path)
    conn_src.backup(conn_dst)
    conn_dst.close()
    
    # 2. Extract Collections & Build Master Excel
    wb = openpyxl.Workbook()
    wb.remove(wb.active) # Remove default sheet
    
    cur = conn_src.cursor()
    cur.execute("SELECT collection, id, data FROM generic_documents ORDER BY collection, id")
    rows = cur.fetchall()
    
    collections_data = {}
    for col_name, doc_id, data_str in rows:
        if col_name not in collections_data:
            collections_data[col_name] = []
        try:
            d = json.loads(data_str)
            d['id'] = doc_id
            collections_data[col_name].append(d)
        except Exception:
            pass
            
    # Save JSON files per collection
    for col_name, items in collections_data.items():
        json_path = os.path.join(json_dir, f"{col_name}.json")
        with open(json_path, 'w', encoding='utf-8') as f:
            json.dump(items, f, ensure_ascii=False, indent=2)
            
    # Build Vault Map by Client ID
    vault = collections_data.get('passwords_vault', [])
    vault_by_client = {}
    for v in vault:
        cid = v.get('clientId', '')
        if cid:
            if cid not in vault_by_client:
                vault_by_client[cid] = []
            info = f"{v.get('title','')}: {v.get('username','')} ({v.get('url','')})"
            vault_by_client[cid].append(info)

    # --- Clients ---
    clients = collections_data.get('clients', [])
    client_headers = ['ID', 'שם', 'איש קשר', 'טלפון', 'אימייל', 'מחלקה', 'סטטוס', 'ח.פ/ע.מ', 'יתרה', 'גישות וכספת מקושרת', 'סוכן', 'נוצר בתאריך']
    client_rows = []
    for c in clients:
        cid = c.get('id', '')
        v_info = " | ".join(vault_by_client.get(cid, []))
        client_rows.append([
            cid, c.get('name', ''), c.get('contact', ''), c.get('phone', ''),
            c.get('email', ''), c.get('department', ''), c.get('status', ''), c.get('taxId', ''),
            c.get('financeBalance', 0), v_info, c.get('salesAgentUserId', ''), c.get('createdAt', '')
        ])
    create_styled_sheet(wb, "01_לקוחות (Clients)", client_headers, client_rows, "1E3A8A")
    
    # --- Leads ---
    leads = collections_data.get('leads', [])
    lead_headers = ['ID', 'שם', 'חברה', 'טלפון', 'אימייל', 'סטטוס', 'צורכי לקוח', 'מקור', 'סוכן', 'מעקב הבא', 'נוצר בתאריך']
    lead_rows = []
    for l in leads:
        lead_rows.append([
            l.get('id', ''), l.get('name', ''), l.get('company', ''), l.get('phone', ''),
            l.get('email', ''), l.get('status', ''), l.get('clientNeeds', ''), l.get('source', ''),
            l.get('salesAgentUserId', ''), l.get('nextFollowUpDate', ''), l.get('createdAt', '')
        ])
    create_styled_sheet(wb, "02_לידים (Leads)", lead_headers, lead_rows, "C2410C")
    
    # --- Tasks ---
    tasks = collections_data.get('tasks', [])
    task_headers = ['ID', 'כותרת', 'תיאור', 'סטטוס', 'עדיפות', 'אחראי', 'לקוח', 'פרויקט', 'תאריך יעד', 'נוצר בתאריך']
    task_rows = []
    for t in tasks:
        task_rows.append([
            t.get('id', ''), t.get('title', ''), t.get('description', ''), t.get('status', ''),
            t.get('priority', ''), t.get('assigneeId', ''), t.get('clientId', ''), t.get('projectId', ''),
            t.get('dueDate', ''), t.get('createdAt', '')
        ])
    create_styled_sheet(wb, "03_משימות (Tasks)", task_headers, task_rows, "047857")

    # --- Projects ---
    projects = collections_data.get('projects', [])
    proj_headers = ['ID', 'שם פרויקט', 'לקוח', 'סטטוס', 'סוג החיוב', 'עדיפות', 'התקדמות %', 'גישות מקושרות ללקוח', 'תאריך יעד', 'נוצר בתאריך']
    proj_rows = []
    for p in projects:
        cid = p.get('clientId', '')
        v_info = " | ".join(vault_by_client.get(cid, []))
        proj_rows.append([
            p.get('id', ''), p.get('name', ''), cid, p.get('status', ''),
            p.get('projectBillingType', ''), p.get('priority', ''), p.get('manualProgressPercent', ''),
            v_info, p.get('dueDate', ''), p.get('createdAt', '')
        ])
    create_styled_sheet(wb, "04_פרויקטים (Projects)", proj_headers, proj_rows, "6D28D9")

    # --- Meetings ---
    meetings = collections_data.get('meetings', [])
    m_headers = ['ID', 'כותרת', 'מטרה', 'סדר יום', 'תאריך', 'שעה', 'לקוח', 'פרויקט', 'סיכום סגירה']
    m_rows = []
    for m in meetings:
        m_rows.append([
            m.get('id', ''), m.get('title', ''), m.get('purpose', ''), m.get('agenda', ''),
            m.get('startDate', ''), m.get('startTime', ''), m.get('clientId', ''), m.get('projectId', ''),
            m.get('completionSummary', '')
        ])
    create_styled_sheet(wb, "05_פגישות (Meetings)", m_headers, m_rows, "B45309")

    # --- Documents ---
    docs = collections_data.get('documents', [])
    doc_headers = ['ID', 'כותרת', 'סוג', 'תוכן / הערות', 'לקוח', 'פרויקט', 'תאריך']
    doc_rows = []
    for d in docs:
        doc_rows.append([
            d.get('id', ''), d.get('title', ''), d.get('type', ''), d.get('content', d.get('notes', '')),
            d.get('clientId', ''), d.get('projectId', ''), d.get('createdAt', '')
        ])
    create_styled_sheet(wb, "06_מסמכים (Docs)", doc_headers, doc_rows, "374151")

    # --- Incomes (Morning) ---
    incomes = collections_data.get('financeIncomes', [])
    inc_headers = ['ID', 'מזהה Morning', 'מספר חשבונית', 'תאריך', 'לקוח מקושר', 'סוג מסמך', 'סכום', 'מטבע', 'סטטוס', 'תיאור']
    inc_rows = []
    for i in incomes:
        inc_rows.append([
            i.get('id', ''), i.get('greenInvoiceId', ''), i.get('invoiceNumber', i.get('greenInvoiceNumber', '')),
            i.get('invoiceDate', ''), i.get('clientId', i.get('greenInvoiceClientName', '')),
            i.get('documentKind', i.get('greenInvoiceType', '')), i.get('amount', ''), i.get('currency', 'ILS'),
            i.get('status', i.get('greenInvoiceStatus', '')), i.get('description', '')
        ])
    create_styled_sheet(wb, "07_הכנסות (Incomes)", inc_headers, inc_rows, "059669")

    # --- Expenses ---
    expenses = collections_data.get('financeExpenses', [])
    exp_headers = ['ID', 'מספר מסמך', 'תאריך', 'ספק', 'סכום', 'מטבע', 'סטטוס', 'תיאור']
    exp_rows = []
    for e in expenses:
        exp_rows.append([
            e.get('id', ''), e.get('documentNumber', ''), e.get('date', ''), e.get('supplierName', ''),
            e.get('totalAmount', ''), e.get('currency', 'ILS'), e.get('status', ''), e.get('description', '')
        ])
    create_styled_sheet(wb, "08_הוצאות (Expenses)", exp_headers, exp_rows, "DC2626")

    # Save Excel Workbook
    master_excel_path = os.path.join(excel_dir, f"os_hub_master_backup_{timestamp_str}.xlsx")
    wb.save(master_excel_path)
    
    # 3. Copy Morning CSVs if exist
    if os.path.exists(MORNING_DIR):
        for f in os.listdir(MORNING_DIR):
            if f.endswith('.csv'):
                shutil.copy2(os.path.join(MORNING_DIR, f), os.path.join(morning_copy_dir, f))

    # 4. Maintain 'latest' symlink/folder copy
    latest_dir = os.path.join(BACKUPS_BASE, "latest")
    if os.path.exists(latest_dir):
        if os.path.islink(latest_dir):
            os.unlink(latest_dir)
        else:
            shutil.rmtree(latest_dir)
    shutil.copytree(target_dir, latest_dir)
    
    conn_src.close()
    
    summary = {
        "status": "success",
        "timestamp": timestamp_str,
        "backup_path": target_dir,
        "master_excel": master_excel_path,
        "collections_count": len(collections_data),
        "total_clients": len(clients),
        "total_leads": len(leads),
        "total_tasks": len(tasks),
        "total_projects": len(projects),
        "total_incomes": len(incomes),
        "total_expenses": len(expenses)
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return summary

if __name__ == '__main__':
    run_backup()
